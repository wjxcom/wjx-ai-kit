from __future__ import annotations

import math
import re
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from .exceptions import InputFileError, ParameterError
from .models import InputRecord

_COLUMN_RE = re.compile(r"^[A-Za-z]+$")


def column_to_index(value: str | int) -> int:
    """Convert a 1-based Excel column number or letters to a 0-based index."""
    text = str(value).strip()
    if not text:
        raise ParameterError("列参数不能为空")
    if text.isdigit():
        number = int(text)
        if number < 1:
            raise ParameterError(f"列号必须从 1 开始: {value}")
        return number - 1
    if not _COLUMN_RE.fullmatch(text):
        raise ParameterError(f"无效的 Excel 列参数: {value}")
    result = 0
    for char in text.upper():
        result = result * 26 + ord(char) - ord("A") + 1
    return result - 1


def _format_value(value: Any, number_format: str | None = None) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if not math.isfinite(value):
            return str(value)
        if value.is_integer():
            return str(int(value))
        return format(value, ".15g")
    text = str(value)
    return text.strip() if isinstance(value, str) else text


def _select_sheet_names_xlsx(path: Path) -> list[str]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            return list(workbook.sheetnames)
        finally:
            workbook.close()
    except Exception as exc:
        raise InputFileError(f"无法读取 Excel 文件: {exc}") from exc


def get_sheet_names(path: Path) -> list[str]:
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        return _select_sheet_names_xlsx(path)
    if suffix == ".xls":
        try:
            import xlrd

            workbook = xlrd.open_workbook(path, on_demand=True)
            try:
                return workbook.sheet_names()
            finally:
                workbook.release_resources()
        except ImportError as exc:
            raise InputFileError("读取 .xls 需要安装 xlrd") from exc
        except Exception as exc:
            raise InputFileError(f"无法读取 Excel 文件: {exc}") from exc
    raise InputFileError("输入文件必须是 .xlsx 或 .xls")


def _resolve_sheet(names: list[str], sheet: str | None, sheet_index: int | None) -> str:
    if sheet is not None and sheet_index is not None:
        raise ParameterError("--sheet 和 --sheet-index 不能同时使用")
    if sheet is not None:
        if sheet not in names:
            raise ParameterError(f"工作表不存在: {sheet}")
        return sheet
    if sheet_index is not None:
        if sheet_index < 1 or sheet_index > len(names):
            raise ParameterError(f"工作表序号超出范围: {sheet_index}")
        return names[sheet_index - 1]
    if not names:
        raise InputFileError("输入文件没有工作表")
    return names[0]


def _is_empty_row(values: list[str]) -> bool:
    return all(not value.strip() for value in values)


def _load_xlsx(path: Path, sheet_name: str, id_index: int, url_index: int) -> list[InputRecord]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
        cached = load_workbook(path, read_only=True, data_only=True)
        try:
            worksheet = workbook[sheet_name]
            cached_sheet = cached[sheet_name]
            max_column = max(worksheet.max_column, 1)
            if id_index >= max_column or url_index >= max_column:
                raise ParameterError("ID 列或 URL 列超出工作表范围")
            records: list[InputRecord] = []
            for row_number, cells in enumerate(worksheet.iter_rows(min_row=2), start=2):
                raw_values = [_format_value(cell.value, cell.number_format) for cell in cells]
                if _is_empty_row(raw_values):
                    continue
                id_cell = cells[id_index]
                url_cell = cells[url_index]
                if id_cell.data_type == "f":
                    cached_id = cached_sheet.cell(row=row_number, column=id_index + 1).value
                    raw_id = _format_value(cached_id if cached_id is not None else id_cell.value)
                else:
                    raw_id = _format_value(id_cell.value, id_cell.number_format)
                if url_cell.data_type == "f":
                    cached_url = cached_sheet.cell(row=row_number, column=url_index + 1).value
                    raw_url = _format_value(cached_url if cached_url is not None else url_cell.value)
                else:
                    raw_url = _format_value(url_cell.value, url_cell.number_format)
                records.append(InputRecord(row_number, raw_id.strip(), raw_url.strip()))
            return records
        finally:
            workbook.close()
            cached.close()
    except (ParameterError, InputFileError):
        raise
    except Exception as exc:
        raise InputFileError(f"无法读取 Excel 数据: {exc}") from exc


def _xls_cell_text(sheet: Any, row: int, col: int, datemode: int) -> str:
    import xlrd

    cell = sheet.cell(row, col)
    if cell.ctype == xlrd.XL_CELL_EMPTY:
        return ""
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return "TRUE" if cell.value else "FALSE"
    if cell.ctype == xlrd.XL_CELL_DATE:
        value = xlrd.xldate_as_datetime(cell.value, datemode)
        return _format_value(value)
    return _format_value(cell.value)


def _load_xls(path: Path, sheet_name: str, id_index: int, url_index: int) -> list[InputRecord]:
    try:
        import xlrd

        workbook = xlrd.open_workbook(path, on_demand=True)
        try:
            sheet = workbook.sheet_by_name(sheet_name)
            if id_index >= sheet.ncols or url_index >= sheet.ncols:
                raise ParameterError("ID 列或 URL 列超出工作表范围")
            records: list[InputRecord] = []
            for row_number in range(1, sheet.nrows):
                values = [_xls_cell_text(sheet, row_number, col, workbook.datemode) for col in range(sheet.ncols)]
                if _is_empty_row(values):
                    continue
                records.append(
                    InputRecord(
                        row_number + 1,
                        values[id_index].strip(),
                        values[url_index].strip(),
                    )
                )
            return records
        finally:
            workbook.release_resources()
    except (ParameterError, InputFileError):
        raise
    except ImportError as exc:
        raise InputFileError("读取 .xls 需要安装 xlrd") from exc
    except Exception as exc:
        raise InputFileError(f"无法读取 Excel 数据: {exc}") from exc


def load_records(path: Path, sheet: str | None, sheet_index: int | None, id_column: str, url_column: str) -> list[InputRecord]:
    if not path.exists():
        raise InputFileError(f"输入文件不存在: {path}")
    if not path.is_file():
        raise InputFileError(f"输入路径不是文件: {path}")
    id_index = column_to_index(id_column)
    url_index = column_to_index(url_column)
    if id_index == url_index:
        raise ParameterError("ID 列和 URL 列不能相同")
    names = get_sheet_names(path)
    selected = _resolve_sheet(names, sheet, sheet_index)
    if path.suffix.lower() == ".xlsx":
        return _load_xlsx(path, selected, id_index, url_index)
    if path.suffix.lower() == ".xls":
        return _load_xls(path, selected, id_index, url_index)
    raise InputFileError("输入文件必须是 .xlsx 或 .xls")

