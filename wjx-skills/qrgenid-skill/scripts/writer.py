from __future__ import annotations

import os
import tempfile
import zipfile
from pathlib import Path

import xlsxwriter
from openpyxl import load_workbook

from .exceptions import OutputError
from .models import InputRecord

HEADERS = ["ID", "URL", "二维码图片", "状态", "失败原因"]


def _write_workbook(path: Path, records: list[InputRecord]) -> None:
    workbook = xlsxwriter.Workbook(str(path))
    try:
        worksheet = workbook.add_worksheet("Sheet1")
        worksheet.set_column("A:A", 22)
        worksheet.set_column("B:B", 58)
        worksheet.set_column("C:C", 22)
        worksheet.set_column("D:D", 12)
        worksheet.set_column("E:E", 42)
        worksheet.write_row(0, 0, HEADERS)
        for output_row, record in enumerate(records, start=1):
            worksheet.write(output_row, 0, record.record_id)
            worksheet.write(output_row, 1, record.url)
            worksheet.write(output_row, 3, record.status)
            worksheet.write(output_row, 4, record.failure_reason)
            if record.status == "成功" and record.image_path and record.image_path.exists():
                from PIL import Image

                with Image.open(record.image_path) as image:
                    width, height = image.size
                scale = min(1.0, 150 / width, 150 / height)
                worksheet.set_row(output_row, max(45, height * scale * 0.75 + 6))
                worksheet.insert_image(
                    output_row,
                    2,
                    str(record.image_path),
                    {"x_scale": scale, "y_scale": scale, "x_offset": 0, "y_offset": 0, "object_position": 1},
                )
            else:
                worksheet.set_row(output_row, 20)
    finally:
        workbook.close()


def verify_workbook(path: Path, expected_records: int, expect_images: int) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            media = [name for name in archive.namelist() if name.startswith("xl/media/")]
            if len(media) != expect_images:
                raise OutputError(f"Excel 图片数量校验失败，预期 {expect_images}，实际 {len(media)}")
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            if workbook.sheetnames != ["Sheet1"]:
                raise OutputError("输出工作表名称校验失败")
            worksheet = workbook["Sheet1"]
            if worksheet.max_row != expected_records + 1 or worksheet.max_column != len(HEADERS):
                raise OutputError("输出 Excel 行列结构校验失败")
        finally:
            workbook.close()
    except OutputError:
        raise
    except Exception as exc:
        raise OutputError(f"无法验证输出 Excel: {exc}") from exc


def write_verified_workbook(output_file: Path, records: list[InputRecord]) -> None:
    temp_fd, temp_name = tempfile.mkstemp(prefix=".qrgenid-", suffix=".xlsx", dir=output_file.parent)
    os.close(temp_fd)
    temp_path = Path(temp_name)
    try:
        _write_workbook(temp_path, records)
        image_count = sum(1 for record in records if record.status == "成功" and record.image_path and record.image_path.exists())
        verify_workbook(temp_path, len(records), image_count)
        os.replace(temp_path, output_file)
    except PermissionError as exc:
        raise OutputError(f"输出文件被占用或没有替换权限: {output_file}") from exc
    except OutputError:
        raise
    except Exception as exc:
        raise OutputError(f"写入输出 Excel 失败: {exc}") from exc
    finally:
        if temp_path.exists():
            try:
                temp_path.unlink()
            except OSError:
                pass

