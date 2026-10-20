"""Embed original-resolution HTTP(S) images in a new Excel data snapshot."""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import re
import threading
from concurrent.futures import ThreadPoolExecutor, as_completed
from copy import copy
from io import BytesIO
from pathlib import Path
from urllib.parse import urlsplit

import requests
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.drawing.picture import PictureLocking
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor, SpreadsheetDrawing
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.units import pixels_to_EMU, pixels_to_points
from PIL import Image, ImageOps, UnidentifiedImageError
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

EMPTY = {"", "(空)", "（空）", "无", "-"}
URL_RE = re.compile(r'https?://[^\s<>"\u3000，；、,;]+?(?=https?://|[\s<>"\u3000，；、,;]|$)', re.I)
LOCAL = threading.local()
PADDING = GAP = 5


def sidecar(output, suffix):
    return output.with_name(output.stem + suffix)


def extract_urls(value):
    return [m.group().rstrip("。") for m in URL_RE.finditer(str(value or ""))]


def column_number(value):
    value = str(value).strip().upper()
    number = int(value) if value.isdigit() else column_index_from_string(value)
    if not 1 <= number <= 16384:
        raise ValueError(f"Column out of range: {value}")
    return number


def http_session():
    if not hasattr(LOCAL, "session"):
        client = requests.Session()
        retries = Retry(total=2, backoff_factor=0.5, status_forcelist=[429, 500, 502, 503, 504])
        adapter = HTTPAdapter(max_retries=retries)
        client.mount("https://", adapter)
        client.mount("http://", adapter)
        LOCAL.session = client
    return LOCAL.session


class NonImage(ValueError):
    pass


def check_image(payload):
    if payload.startswith(b"%PDF"):
        raise NonImage("PDF 文档，不是图片；未插入")
    if payload.startswith(b"PK\x03\x04"):
        raise NonImage("ZIP/Office 文档，不是图片；未插入")
    with Image.open(BytesIO(payload)) as image:
        image.verify()


def fetch_image(url, cache, display_width, max_bytes):
    try:
        parts = urlsplit(url)
        if parts.scheme.lower() not in {"http", "https"} or not parts.hostname:
            return {"status": "invalid_url", "reason": "无效 HTTP(S) URL"}
        digest = hashlib.sha256(url.encode("utf-8")).hexdigest()
        original = cache / (digest + ".download")
        if original.exists():
            if original.stat().st_size > max_bytes:
                raise ValueError("缓存文件超过下载大小上限")
            payload = original.read_bytes()
            check_image(payload)
        else:
            with http_session().get(url, timeout=(15, 45), stream=True) as response:
                response.raise_for_status()
                if int(response.headers.get("Content-Length", 0)) > max_bytes:
                    raise ValueError("文件超过下载大小上限")
                chunks, size = [], 0
                for chunk in response.iter_content(1024 * 128):
                    size += len(chunk)
                    if size > max_bytes:
                        raise ValueError("文件超过下载大小上限")
                    chunks.append(chunk)
                payload = b"".join(chunks)
                check_image(payload)
            original.write_bytes(payload)
        with Image.open(BytesIO(payload)) as opened:
            fmt, orientation = opened.format, opened.getexif().get(274, 1)
            image = ImageOps.exif_transpose(opened)
            original_width, original_height = image.size
            display_height = original_height * display_width / original_width
            if fmt in {"JPEG", "PNG"} and orientation == 1:
                extension = ".jpeg" if fmt == "JPEG" else ".png"
                target = cache / (digest + "_original" + extension)
                # Replace only this script's derived cache file, never user inputs.
                if not target.exists() or hashlib.sha256(target.read_bytes()).digest() != hashlib.sha256(payload).digest():
                    target.write_bytes(payload)
            else:
                target = cache / (digest + "_original.png")
                image = image.convert("RGBA" if "A" in image.getbands() else "RGB")
                image.save(target, "PNG")
        return {"status": "success", "path": str(target.resolve()), "original_width_px": original_width,
                "original_height_px": original_height, "height_px": display_height, "format": fmt,
                "sha256": hashlib.sha256(target.read_bytes()).hexdigest()}
    except NonImage as error:
        return {"status": "non_image", "reason": str(error)}
    except (requests.RequestException, OSError) as error:
        status = "invalid_image" if isinstance(error, UnidentifiedImageError) else "download_failed"
        return {"status": status, "reason": f"{type(error).__name__}: {error}"}
    except Exception as error:
        return {"status": "invalid_image", "reason": f"{type(error).__name__}: {error}"}


def run(args):
    source_path = args.source.resolve()
    if source_path.suffix.lower() != ".xlsx":
        raise ValueError("Source must be .xlsx")
    output = (args.output or args.source.with_name(args.source.stem + "_images.xlsx")).resolve()
    if output.suffix.lower() != ".xlsx":
        raise ValueError("Output must be .xlsx")
    if output == source_path:
        raise ValueError("Output must not overwrite the source")
    artifacts = [output, sidecar(output, "_issues.csv"), sidecar(output, "_manifest.json")]
    for path in artifacts:
        if path.exists():
            raise ValueError(f"Output already exists; choose another name: {path}")
    if not output.parent.is_dir():
        raise ValueError("Output parent directory does not exist")
    columns = sorted(set(column_number(c) for c in args.columns))
    if args.keep_columns + len(columns) > 16384:
        raise ValueError("Too many output columns")
    source = load_workbook(source_path)
    cached = load_workbook(source_path, data_only=True)
    try:
        if args.sheets:
            names = list(dict.fromkeys(args.sheets))
            if set(names) - set(source.sheetnames):
                raise ValueError(f"Unknown sheets: {set(names) - set(source.sheetnames)}")
        elif len(source.sheetnames) == 1:
            names = source.sheetnames
        else:
            raise ValueError(f"Multiple sheets: {source.sheetnames}; choose explicitly with --sheets")
        # Validate data before any network calls or output files.
        for name in names:
            sheet = source[name]
            if max(columns) > sheet.max_column or args.keep_columns > sheet.max_column:
                raise ValueError(f"{name}: requested columns exceed source dimensions")
            for merged in sheet.merged_cells.ranges:
                if merged.min_col <= args.keep_columns < merged.max_col:
                    raise ValueError(f"{name}!{merged}: merge crosses copied-column boundary")
            for row in sheet.iter_rows(max_col=args.keep_columns or 1):
                for cell in row[:args.keep_columns]:
                    if cell.data_type == "f" and cached[name][cell.coordinate].value is None:
                        raise ValueError(f"{name}!{cell.coordinate}: formula has no cached value; recalculate and save source in Excel/WPS")
        cache = (args.cache or sidecar(output, "_cache")).resolve()
        cache.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        workbook.remove(workbook.active)
        jobs, issues, sheet_records = [], [], []
        for name in names:
            sheet, dest = source[name], workbook.create_sheet(name)
            dest.freeze_panes = "A2"
            dest.sheet_view.showGridLines = sheet.sheet_view.showGridLines
            sheet_records.append({"name": name, "rows": sheet.max_row})
            for c in range(1, args.keep_columns + 1):
                letter = get_column_letter(c)
                dest.column_dimensions[letter].width = sheet.column_dimensions[letter].width or 15
            for r in range(1, sheet.max_row + 1):
                dest.row_dimensions[r].height = sheet.row_dimensions[r].height or 15
                for c in range(1, args.keep_columns + 1):
                    old, new = sheet.cell(r, c), dest.cell(r, c)
                    value = cached[name].cell(r, c).value if old.data_type == "f" else old.value
                    new.value, new._style = value, copy(old._style)
                    if old.data_type != "f":
                        new.data_type = old.data_type
                    if old.hyperlink:
                        new.hyperlink = copy(old.hyperlink)
                        new.hyperlink.ref = new.coordinate
                    if old.comment:
                        new.comment = copy(old.comment)
            for merged in sheet.merged_cells.ranges:
                if merged.max_col <= args.keep_columns:
                    dest.merge_cells(str(merged))
            for out_col, source_col in enumerate(columns, start=args.keep_columns + 1):
                old, new = sheet.cell(1, source_col), dest.cell(1, out_col)
                new.value, new._style = old.value, copy(old._style)
                alignment = copy(old.alignment)
                alignment.wrap_text = True
                new.alignment = alignment
                dest.column_dimensions[get_column_letter(out_col)].width = 20
                for r in range(2, sheet.max_row + 1):
                    cell, target = sheet.cell(r, source_col), dest.cell(r, out_col)
                    text = str(cell.value if cell.value is not None else "").strip()
                    urls = extract_urls(text)
                    if text not in EMPTY:
                        target.comment = Comment(text, "Source URLs")
                    if not urls and text not in EMPTY:
                        issues.append({"sheet": name, "source": cell.coordinate, "target": target.coordinate,
                                       "slot": None, "url": text, "status": "invalid_url", "reason": "非空单元格未找到 HTTP(S) URL"})
                    if urls:
                        jobs.append((name, r, source_col, out_col, urls))
            dest.auto_filter.ref = f"A1:{get_column_letter(args.keep_columns + len(columns))}{sheet.max_row}"
        urls = list(dict.fromkeys(url for *_, items in jobs for url in items))
        print(f"Processing {len(urls)} unique URLs from {len(jobs)} cells", flush=True)
        results = {}
        with ThreadPoolExecutor(max_workers=args.workers) as pool:
            futures = {pool.submit(fetch_image, url, cache, args.width, int(args.max_download_mb * 1024**2)): url for url in urls}
            for n, future in enumerate(as_completed(futures), 1):
                results[futures[future]] = future.result()
                if n % 25 == 0 or n == len(urls):
                    print(f"Completed {n}/{len(urls)}", flush=True)
        images, cell_records, column_pixels = [], [], {}
        for name, r, source_col, out_col, items in jobs:
            dest = workbook[name]
            target = f"{get_column_letter(out_col)}{r}"
            cell_records.append({"sheet": name, "source": f"{get_column_letter(source_col)}{r}", "target": target, "url_count": len(items)})
            total = 2 * PADDING + len(items) * args.width + (len(items) - 1) * GAP
            column_pixels[(name, out_col)] = max(column_pixels.get((name, out_col), 0), total)
            for slot, url in enumerate(items):
                result = results[url]
                common = {"sheet": name, "source": f"{get_column_letter(source_col)}{r}", "target": target, "slot": slot}
                if result["status"] != "success":
                    issues.append(dict(common, url=url, status=result["status"], reason=result["reason"]))
                    continue
                height = result["height_px"]
                required = pixels_to_points(height + 2 * PADDING)
                if required > 409:
                    raise ValueError(f"{name}!{common['source']}: exceeds Excel row height limit; use smaller --width (original pixels will be preserved)")
                image = ExcelImage(result["path"])
                offset = PADDING + slot * (args.width + GAP)
                marker = AnchorMarker(col=out_col - 1, row=r - 1, colOff=pixels_to_EMU(offset), rowOff=pixels_to_EMU(PADDING))
                image.anchor = OneCellAnchor(_from=marker, ext=XDRPositiveSize2D(pixels_to_EMU(args.width), pixels_to_EMU(height)))
                image.anchor.pic = SpreadsheetDrawing()._picture_frame(len(images) + 1)
                image.anchor.pic.nvPicPr.cNvPicPr.picLocks = PictureLocking(noChangeAspect=True)
                dest.add_image(image)
                dest.row_dimensions[r].height = max(dest.row_dimensions[r].height or 15, required)
                images.append(dict(common, offset_px=offset, width_px=args.width, height_px=height,
                                   original_width_px=result["original_width_px"], original_height_px=result["original_height_px"],
                                   sha256=result["sha256"], format=result["format"]))
        for (name, c), pixels in column_pixels.items():
            width = math.ceil((pixels - 5) / 7 * 100) / 100
            if width > 255:
                raise ValueError(f"{name} column {c}: exceeds Excel width limit; use smaller --width")
            workbook[name].column_dimensions[get_column_letter(c)].width = max(20, width)
        workbook.save(output)
        with sidecar(output, "_issues.csv").open("w", encoding="utf-8-sig", newline="") as stream:
            fields = ["sheet", "source", "target", "slot", "status", "url", "reason"]
            writer = csv.DictWriter(stream, fieldnames=fields)
            writer.writeheader()
            # Prevent untrusted cell text from becoming a spreadsheet formula
            # when a user opens the issue CSV in Excel.
            for issue in issues:
                safe = {key: "'" + value if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")) else value
                        for key, value in issue.items()}
                writer.writerow(safe)
        manifest = {"schema_version": 1, "source": str(source_path), "source_sha256": hashlib.sha256(source_path.read_bytes()).hexdigest(),
                    "output": str(output), "keep_columns": args.keep_columns, "image_columns": columns, "width_px": args.width,
                    "padding_px": PADDING, "gap_px": GAP, "formula_mode": "cached", "sheets": sheet_records,
                    "preserve_original_resolution": True, "embedded_images": len(images), "unique_urls": len(urls),
                    "issue_count": len(issues), "cells": cell_records, "images": images, "issues": issues}
        sidecar(output, "_manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
        print(json.dumps({"output": str(output), "embedded_images": len(images), "issues": len(issues),
                          "size_mib": round(output.stat().st_size / 1024**2, 2)}, ensure_ascii=False), flush=True)
        return manifest
    finally:
        source.close()
        cached.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("--columns", nargs="+", required=True)
    parser.add_argument("--keep-columns", type=int, default=6)
    parser.add_argument("--width", type=float, default=100)
    parser.add_argument("--sheets", nargs="+")
    parser.add_argument("--output", type=Path)
    parser.add_argument("--cache", type=Path)
    parser.add_argument("--workers", type=int, default=4)
    parser.add_argument("--max-download-mb", type=float, default=50)
    args = parser.parse_args()
    if not (math.isfinite(args.width) and args.width > 0 and args.workers > 0
            and 0 <= args.keep_columns <= 16384 and math.isfinite(args.max_download_mb) and args.max_download_mb > 0):
        parser.error("Width, workers and maximum download size must be positive; keep-columns must be 0..16384")
    try:
        run(args)
    except Exception as error:
        parser.exit(1, f"Error: {error}\n")


if __name__ == "__main__":
    main()
