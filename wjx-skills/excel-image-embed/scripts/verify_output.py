"""Verify embedded image bytes, resolution, proportional layout and source values."""
from __future__ import annotations

import argparse
import hashlib
import json
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU, pixels_to_points
from PIL import Image
from io import BytesIO


def verify(output, manifest_path=None):
    output = Path(output)
    manifest_path = Path(manifest_path) if manifest_path else output.with_name(output.stem + "_manifest.json")
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    checks, warnings = [], []

    def check(name, condition, detail=""):
        checks.append({"name": name, "passed": bool(condition), "detail": detail})

    with zipfile.ZipFile(output) as archive:
        check("zip_integrity", archive.testzip() is None)
        for member in archive.namelist():
            if member.endswith((".xml", ".rels")):
                ET.fromstring(archive.read(member))
        check("xml_parts_parse", True)
        check("media_count", len([p for p in archive.namelist() if p.startswith("xl/media/")]) == manifest["embedded_images"])
    workbook = load_workbook(output)
    try:
        check("sheet_selection", workbook.sheetnames == [s["name"] for s in manifest["sheets"]])
        check("image_count", sum(len(s._images) for s in workbook) == manifest["embedded_images"] == len(manifest["images"]))
        check("issue_count", len(manifest["issues"]) == manifest["issue_count"])
        for sheet in workbook:
            expected = [i for i in manifest["images"] if i["sheet"] == sheet.title]
            check(f"{sheet.title}:image_count", len(expected) == len(sheet._images))
            for index, (image, item) in enumerate(zip(sheet._images, expected)):
                label = f"{sheet.title}!{item['target']}[{item['slot']}]#{index}"
                marker, ext = image.anchor._from, image.anchor.ext
                target = sheet[item["target"]]
                check(label + ":source_row", int(''.join(filter(str.isdigit, item["source"]))) == target.row)
                check(label + ":anchor", marker.row + 1 == target.row and marker.col + 1 == target.column)
                expected_offset = manifest["padding_px"] + item["slot"] * (manifest["width_px"] + manifest["gap_px"])
                check(label + ":slot_offset", marker.colOff == pixels_to_EMU(expected_offset) == pixels_to_EMU(item["offset_px"]))
                check(label + ":vertical_padding", marker.rowOff == pixels_to_EMU(manifest["padding_px"]))
                check(label + ":display_width", ext.cx == pixels_to_EMU(manifest["width_px"]))
                check(label + ":proportional_height", ext.cy == pixels_to_EMU(item["original_height_px"] * manifest["width_px"] / item["original_width_px"]))
                lock = image.anchor.pic.nvPicPr.cNvPicPr.picLocks
                check(label + ":aspect_ratio_lock", lock is not None and lock.noChangeAspect is True)
                payload = image._data()
                check(label + ":original_file_hash", hashlib.sha256(payload).hexdigest() == item["sha256"])
                with Image.open(BytesIO(payload)) as raster:
                    check(label + ":original_pixel_dimensions", raster.size == (item["original_width_px"], item["original_height_px"]))
                    raster.verify()
                height = pixels_to_points(item["height_px"] + 2 * manifest["padding_px"])
                actual_height = sheet.row_dimensions[target.row].height or 15
                check(label + ":row_height", actual_height + 1e-7 >= height and actual_height <= 409)
                pixels = sheet.column_dimensions[target.column_letter].width * 7 + 5
                check(label + ":within_column", pixels + 1e-7 >= item["offset_px"] + manifest["width_px"] + manifest["padding_px"])
            for cell in (c for c in manifest["cells"] if c["sheet"] == sheet.title):
                images = [i for i in expected if i["target"] == cell["target"]]
                issues = [i for i in manifest["issues"] if i["sheet"] == sheet.title and i["target"] == cell["target"] and i["slot"] is not None]
                slots = [i["slot"] for i in images + issues]
                check(f"{sheet.title}!{cell['target']}:all_url_slots_accounted", sorted(slots) == list(range(cell["url_count"])))
                required = 2 * manifest["padding_px"] + cell["url_count"] * manifest["width_px"] + (cell["url_count"] - 1) * manifest["gap_px"]
                col = sheet[cell["target"]].column_letter
                check(f"{sheet.title}!{cell['target']}:total_column_width", sheet.column_dimensions[col].width * 7 + 5 + 1e-7 >= required)
        source_path = Path(manifest["source"])
        if source_path.exists():
            check("source_unchanged", hashlib.sha256(source_path.read_bytes()).hexdigest() == manifest["source_sha256"])
            source = load_workbook(source_path, data_only=True)
            raw = load_workbook(source_path)
            try:
                for record in manifest["sheets"]:
                    sheet, old = workbook[record["name"]], source[record["name"]]
                    check(f"{sheet.title}:source_rows", max(sheet.max_row, max(sheet.row_dimensions, default=0)) == record["rows"])
                    for r in range(1, record["rows"] + 1):
                        for c in range(1, manifest["keep_columns"] + 1):
                            check(f"{sheet.title}!{get_column_letter(c)}{r}:source_value", sheet.cell(r, c).value == old.cell(r, c).value)
                    for out_col, source_col in enumerate(manifest["image_columns"], start=manifest["keep_columns"] + 1):
                        check(f"{sheet.title}:{out_col}:header", sheet.cell(1, out_col).value == raw[record["name"]].cell(1, source_col).value)
            finally:
                source.close()
                raw.close()
        else:
            warnings.append("Source is unavailable at its original path; source-data comparison skipped. Embedded images and layout were still checked.")
    finally:
        workbook.close()
    return {"passed": all(c["passed"] for c in checks), "checks": checks, "warnings": warnings,
            "summary": {"total": len(checks), "passed": sum(c["passed"] for c in checks), "failed": sum(not c["passed"] for c in checks)}}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("output", type=Path)
    parser.add_argument("--manifest", type=Path)
    args = parser.parse_args()
    report_path = args.output.with_name(args.output.stem + "_verification.json")
    try:
        report = verify(args.output, args.manifest)
    except Exception as error:
        report = {"passed": False, "error": f"{type(error).__name__}: {error}"}
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({k: v for k, v in report.items() if k != "checks"}, ensure_ascii=False))
    if not report["passed"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
