"""Portable, synthetic-only regression tests; no public network or business data."""
from __future__ import annotations

import argparse
import hashlib
import json
import subprocess
import sys
import tempfile
import threading
from functools import partial
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

from openpyxl import Workbook, load_workbook
from PIL import Image, ImageDraw, ImageOps

from embed_images import extract_urls
from verify_output import verify


class QuietHandler(SimpleHTTPRequestHandler):
    def log_message(self, *_):
        pass


def make_fixtures(root, base_url):
    root = Path(root)
    root.mkdir(parents=True, exist_ok=True)
    files = root / "files"
    files.mkdir(exist_ok=True)
    specs = [("portrait.png", (1200, 2400), "#e8f4ff"), ("landscape.jpeg", (1800, 600), "#ffefcd"),
             ("large.png", (2400, 3600), "#d7ffdd"), ("sample.webp", (1500, 750), "#efd7ff")]
    for name, size, color in specs:
        image = Image.new("RGB", size, color)
        draw = ImageDraw.Draw(image)
        for x in range(0, size[0], 50):
            draw.line((x, 0, x, size[1] - 1), fill="black")
        for y in range(0, size[1], 50):
            draw.line((0, y, size[0] - 1, y), fill="black")
        draw.text((20, 20), "SYNTHETIC ONLY - original pixels retained", fill="red")
        image.save(files / name, lossless=True) if name.endswith(".webp") else image.save(files / name)
    rotated = Image.new("RGB", (900, 1500), "#ffe2e2")
    exif = Image.Exif()
    exif[274] = 6
    rotated.save(files / "rotated.jpeg", exif=exif)
    (files / "document.pdf").write_bytes(b"%PDF-1.4\nSynthetic test, not an image\n%%EOF")
    import zipfile
    with zipfile.ZipFile(files / "document.docx", "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
    (files / "broken.png").write_bytes(b"not an image")
    url = lambda name: base_url.rstrip("/") + "/files/" + name
    inputs = {
        "normal": [
            [url("portrait.png") + "，" + url("landscape.jpeg"), url("landscape.jpeg")],
            [url("landscape.jpeg"), url("portrait.png")],
            [None, url("portrait.png") + ";" + url("landscape.jpeg")],
        ],
        "failures": [
            [None, "(空)"],
            [url("portrait.png") + "," + url("missing.png") + "，" + url("landscape.jpeg"), None],
            ["not-a-valid-url", url("document.pdf")],
            [url("document.docx"), url("broken.png")],
            [url("portrait.png"), url("landscape.jpeg")],
        ],
        "quality": [
            [url("sample.webp"), url("large.png")],
            [url("rotated.jpeg"), url("landscape.jpeg")],
        ],
    }
    expectations = {}
    for name, rows in inputs.items():
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Data"
        sheet.append(["ID", "Name", "Number", "Flag", "Notes", "Code", "Photos G", "Photos H"])
        images, issues = [], []
        width = 125 if name == "quality" else 100
        for r, (g, h) in enumerate(rows, 2):
            sheet.append([r - 1, f"Synthetic-{r}", 1.25 * r, bool(r % 2), "test only", f"ID-{r:03d}", g, h])
            for c, text in ((7, g), (8, h)):
                urls = extract_urls(text)
                if text not in (None, "", "(空)") and not urls:
                    issues.append({"source": f"{'G' if c == 7 else 'H'}{r}", "status": "invalid_url"})
                for slot, item in enumerate(urls):
                    file = files / item.rsplit("/", 1)[-1]
                    if not file.exists() or file.suffix in {".pdf", ".docx"} or file.name == "broken.png":
                        status = "download_failed" if not file.exists() else "invalid_image" if file.name == "broken.png" else "non_image"
                        issues.append({"source": f"{'G' if c == 7 else 'H'}{r}", "slot": slot, "status": status})
                        continue
                    with Image.open(file) as opened:
                        corrected = ImageOps.exif_transpose(opened)
                        images.append({"row": r, "column": c, "slot": slot, "width": width,
                                       "pixels": list(corrected.size), "sha256": hashlib.sha256(file.read_bytes()).hexdigest()
                                       if opened.format in {"PNG", "JPEG"} and opened.getexif().get(274, 1) == 1 else None,
                                       "format": "png" if opened.format == "WEBP" else None})
        workbook.save(root / (name + ".xlsx"))
        expectations[name] = {"images": images, "issues": issues, "width": width}
    (root / "expected.json").write_text(json.dumps(expectations, indent=2), encoding="utf-8")
    return expectations


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path)
    args = parser.parse_args()
    root = args.output_dir or Path(tempfile.mkdtemp(prefix="excel-image-embed-test-"))
    if args.output_dir:
        if root.exists():
            parser.error("output-dir already exists; choose a new directory")
        root.mkdir(parents=True)
    handler = partial(QuietHandler, directory=str(root.resolve()))
    server = ThreadingHTTPServer(("127.0.0.1", 0), handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    results = []
    try:
        expected = make_fixtures(root, f"http://127.0.0.1:{server.server_port}")
        script = Path(__file__).with_name("embed_images.py")
        for name, specification in expected.items():
            source = root / (name + ".xlsx")
            source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
            output = root / (name + "_images.xlsx")
            command = [sys.executable, str(script), str(source), "--columns", "H", "G", "--width", str(specification["width"]), "--output", str(output)]
            subprocess.run(command, check=True)
            report = verify(output)
            assert report["passed"], [c for c in report["checks"] if not c["passed"]]
            manifest = json.loads(output.with_name(output.stem + "_manifest.json").read_text(encoding="utf-8"))
            assert len(manifest["images"]) == len(specification["images"])
            assert len(manifest["issues"]) == len(specification["issues"])
            for actual, wanted in zip(manifest["images"], sorted(specification["images"], key=lambda i: (i["column"], i["row"], i["slot"]))):
                assert [actual["original_width_px"], actual["original_height_px"]] == wanted["pixels"]
                if wanted["sha256"]:
                    assert actual["sha256"] == wanted["sha256"]
            assert hashlib.sha256(source.read_bytes()).hexdigest() == source_hash
            # Existing outputs and source paths must never be overwritten.
            assert subprocess.run(command, capture_output=True).returncode != 0
            assert subprocess.run([sys.executable, str(script), str(source), "--columns", "G", "--output", str(source)], capture_output=True).returncode != 0
            results.append({"case": name, "passed": True, "images": len(manifest["images"]), "issues": len(manifest["issues"]), "checks": report["summary"]["total"]})
        # Multi-sheet ambiguity and zero copied columns.
        workbook = load_workbook(root / "normal.xlsx")
        workbook.create_sheet("Second")
        multi = root / "multiple.xlsx"
        workbook.save(multi)
        rejected = subprocess.run([sys.executable, str(script), str(multi), "--columns", "G"], capture_output=True)
        assert rejected.returncode != 0 and b"Multiple sheets" in rejected.stderr
        selected = root / "selected_images.xlsx"
        subprocess.run([sys.executable, str(script), str(multi), "--columns", "7", "--sheets", "Data", "--keep-columns", "0", "--output", str(selected)], check=True)
        assert verify(selected)["passed"]
        # Reject partial merges and uncached formulas before downloading.
        for kind in ("merge", "formula"):
            workbook = load_workbook(root / "normal.xlsx")
            if kind == "merge":
                workbook.active.merge_cells("F2:G2")
                message = b"merge crosses"
            else:
                workbook.active["C2"] = "=1+2"
                message = b"formula has no cached"
            input_path = root / (kind + ".xlsx")
            workbook.save(input_path)
            rejected = subprocess.run([sys.executable, str(script), str(input_path), "--columns", "G", "H"], capture_output=True)
            assert rejected.returncode != 0 and message in rejected.stderr
            assert not input_path.with_name(input_path.stem + "_images.xlsx").exists()
        # Verify the verifier detects a tampered anchor rather than always passing.
        import shutil
        original = root / "quality_images.xlsx"
        tampered = root / "tampered.xlsx"
        workbook = load_workbook(original)
        workbook.active._images[0].anchor.ext.cx *= 2
        workbook.save(tampered)
        shutil.copyfile(original.with_name(original.stem + "_manifest.json"), tampered.with_name(tampered.stem + "_manifest.json"))
        assert not verify(tampered)["passed"]
        print(json.dumps({"passed": True, "cases": results, "extra_checks": ["no-overwrite", "multi-sheet-confirmation", "numeric-columns", "zero-copied-columns", "partial-merge-rejected", "uncached-formula-rejected", "tampered-anchor-detected"], "directory": str(root)}, indent=2))
        (root / "self-test-results.json").write_text(json.dumps(results, indent=2), encoding="utf-8")
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)


if __name__ == "__main__":
    main()
